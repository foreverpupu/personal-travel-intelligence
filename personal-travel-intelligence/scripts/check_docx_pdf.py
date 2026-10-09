#!/usr/bin/env python3
"""DOCX/PDF smoke QA; page bounds are failures, narrow margins and extra text need review."""
import collections, json, re, sys
try:
    from docx import Document
    import fitz
except ImportError as exc:
    print(f"Missing dependency: {exc}. Install: python3 -m pip install pymupdf python-docx", file=sys.stderr)
    sys.exit(2)
if (len(sys.argv) not in (3, 4, 5)
    or any(a.startswith("-") and a != "--require-chinese" for a in sys.argv[1:])
    or (len(sys.argv)==5 and sys.argv[4] != "--require-chinese")
    or (len(sys.argv)==5 and sys.argv[3] == "--require-chinese")):
    print("Usage: check_docx_pdf.py file.docx file.pdf [report.json] [--require-chinese]", file=sys.stderr)
    sys.exit(2)
from pathlib import Path
if not all(Path(p).is_file() for p in sys.argv[1:3]):
    print("Input DOCX/PDF file not found", file=sys.stderr)
    sys.exit(2)

def cell_text(cell):
    items=[p.text for p in cell.paragraphs]
    for nested in cell.tables:
        for row in nested.rows:
            for c in row.cells: items.extend(cell_text(c))
    return items

def all_word_text(doc):
    items=[]
    # iterate body XML order to include tables, without double-counting tables
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    for node in doc.element.body.iterchildren():
        if node.tag==qn('w:p'): items.append(Paragraph(node,doc).text)
        elif node.tag==qn('w:tbl'):
            t=Table(node,doc)
            for row in t.rows:
                for c in row.cells: items.extend(cell_text(c))
    # Header/footer not added to expected body text because PDF repeats them on each page;
    # explicitly list them as explained extra-text sources.
    return '\n'.join(items)

word=Document(sys.argv[1]); pdf=fitz.open(sys.argv[2])
src=all_word_text(word)
page_text=[p.get_text(sort=True) for p in pdf]
observed='\n'.join(page_text)
han=lambda t:collections.Counter(re.findall(r'[\u3400-\u9fff]',t))
# Latin word comparison is case-insensitive and preserves repeats; matching is not
# semantic translation checking. Headers/footers may increase PDF extra words.
latin=lambda t:collections.Counter(w.lower() for w in re.findall(r'[A-Za-z]+(?:[’\'-][A-Za-z]+)*',t))
missing_latin=latin(src)-latin(observed)
missing=han(src)-han(observed)
extra=han(observed)-han(src)
outside=[]; margin=[]
for page_num,page in enumerate(pdf,1):
    for block in page.get_text('dict')['blocks']:
        for line in block.get('lines',[]):
            for span in line.get('spans',[]):
                x0,y0,x1,y1=span['bbox']
                if x0 < -0.5 or x1 > page.rect.width+.5 or y0 < -.5 or y1>page.rect.height+.5:
                    outside.append([page_num,span['text'][:60]])
                if x0 < 20 or x1 > page.rect.width-20:
                    margin.append([page_num,span['text'][:60]])
# Extra text may be headers/footers repeated on every page; report rather than silently accepting.
headfoot=[]
for section in word.sections:
    headfoot.extend(p.text for p in section.header.paragraphs)
    headfoot.extend(p.text for p in section.footer.paragraphs)
report={'pages':len(pdf),'page_sizes':[list(map(lambda x:round(x,1),(p.rect.width,p.rect.height))) for p in pdf], 'docx_chinese_chars':sum(han(src).values()),'pdf_chinese_chars':sum(han(observed).values()),'missing_chinese':sum(missing.values()),'missing_details':dict(missing),'docx_latin_words':sum(latin(src).values()),'missing_latin_words':sum(missing_latin.values()),'missing_latin_details':dict(missing_latin),'extra_chinese':sum(extra.values()),'extra_details':dict(extra),'possible_header_footer_sources':[x for x in headfoot if x.strip()],'outside_page_spans':outside,'near_margin_spans':margin,'table_count':len(word.tables)}
a4_failures=[{'page':i+1,'size':[round(page.rect.width,1),round(page.rect.height,1)]} for i,page in enumerate(pdf) if not (abs(page.rect.width-595.28)<=2 and abs(page.rect.height-841.89)<=2)]
report['non_a4_pages']=a4_failures
report['no_chinese_in_docx']=sum(han(src).values())==0
report['no_chinese_note']='English-only documents do not exercise Chinese glyph QA' if report['no_chinese_in_docx'] else None
 # PDF headers and footers may legitimately add text; report extra text for manual review.
report['extra_text_requires_review'] = bool(extra or (latin(observed)-latin(src)))
report['near_margin_requires_review'] = bool(margin)
report['docx_empty'] = not src.strip()
report['pdf_text_empty'] = not observed.strip()
report['require_chinese'] = '--require-chinese' in sys.argv
report['failures'] = []
if missing: report['failures'].append('missing_chinese')
if missing_latin: report['failures'].append('missing_latin_words')
if outside: report['failures'].append('outside_page_spans')
if a4_failures: report['failures'].append('non_a4_pages')
if report['docx_empty']: report['failures'].append('empty_docx_body')
if report['pdf_text_empty']: report['failures'].append('empty_pdf_text_layer')
if report['require_chinese'] and report['no_chinese_in_docx']: report['failures'].append('chinese_required_but_absent')
if len(sys.argv)>=4 and sys.argv[3]!='--require-chinese':
    Path(sys.argv[3]).write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
sys.exit(1 if report['failures'] else 0)
